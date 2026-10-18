"""把 state.db 导出成 Excel(多 sheet),供飞书多维表格「导入 Excel」使用。

    python export_excel.py

生成 飞书导入_数据.xlsx(项目根),里面 3 个表:视频池 / 账号 / 发布记录。
在飞书多维表格里点「导入 Excel」选它即可。想更新就再跑一次、重新导入。
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
from src import brain  # noqa: E402

OUT = ROOT / "飞书导入_数据.xlsx"


def main():
    con = brain.connect()
    accts = {r["account_id"]: r["name"]
             for r in con.execute("SELECT account_id,name FROM accounts").fetchall()}

    wb = Workbook()

    ws = wb.active
    ws.title = "视频池"
    ws.append(["视频", "关键词", "序号", "状态", "大小MB", "生成时间"])
    for v in con.execute("SELECT * FROM videos ORDER BY keyword,seq").fetchall():
        ws.append([v["video_id"], v["keyword"], v["seq"], v["status"],
                   v["size_mb"], v["created_at"]])

    ws = wb.create_sheet("账号")
    ws.append(["名称", "account_id", "每日下限", "每日上限", "启用"])
    for r in con.execute("SELECT * FROM accounts").fetchall():
        ws.append([r["name"], r["account_id"], r["daily_min"],
                   r["daily_max"], "是" if r["active"] else "否"])

    ws = wb.create_sheet("发布记录")
    ws.append(["视频", "账号", "关键词", "状态", "排定时间", "task_id", "记录时间"])
    for r in con.execute("SELECT * FROM publish_log ORDER BY id DESC").fetchall():
        ws.append([r["video_id"], accts.get(r["account_id"], r["account_id"]),
                   r["keyword"], r["status"], r["scheduled_time"],
                   r["task_id"], r["created_at"]])

    con.close()
    wb.save(OUT)
    print(f"已导出: {OUT}")


if __name__ == "__main__":
    main()
